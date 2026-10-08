from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'specifikace'
OUT.mkdir(exist_ok=True)
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(9)
doc.core_properties.title = 'Domácí plán a zadání webových her'
doc.core_properties.subject = 'Podpora porozumění řeči a kognitivních dovedností při zohlednění zrakového nálezu'
doc.core_properties.author = 'Klidné hraní'
doc.core_properties.last_modified_by = ''
doc.core_properties.comments = 'Anonymizovaná veřejná specifikace'
doc.core_properties.keywords = ''
doc.core_properties.identifier = ''
footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Domácí plán a webové hry  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(8)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text): doc.add_heading(text, 2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)

def bullets(items):
    for x in items: p(x, 'List Bullet')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c,w in zip(t.columns, widths): c.width = Cm(w)
    for c,txt,w in zip(t.rows[0].cells,headers,widths): c.text=txt; c.width=Cm(w)
    trpr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    for row in rows:
        for c,txt,w in zip(t.add_row().cells,row,widths): c.text=str(txt); c.width=Cm(w)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE7ED' if ri==0 else ('F5F7F8' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1.0
                for run in para.runs:
                    run.font.size=Pt(9.5); run.bold=(ri==0)
    p('')
    return t

def link(text, url):
    para=p(''); rel=para.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    el=OxmlElement('w:hyperlink'); el.set(qn('r:id'), rel)
    run=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'155C83'); pr.append(color); run.append(pr)
    tx=OxmlElement('w:t'); tx.text=text; run.append(tx); el.append(run); para._p.append(el)

p('Domácí plán a zadání webových her', 'Title')
p('Anonymizované funkční východisko  |  Mladší školní věk  |  Verze 1.2', 'Subtitle')
p('Cílem je podpořit porozumění řeči, pracovní paměť a zvládání jednoduchých úkolů při šetrném dávkování práce na blízko. Sada obsahuje osm krátkých her; původní návrh prvních čtyř her dále rozvíjejí hry G5 až G8, společné aktivity s rodičem a přehled skutečně pozorovaných změn. Jde o návrh domácí podpory a vývoje, nikoli předpis oční léčby nebo potvrzení vývojové diagnózy.')
h('Anonymizované východisko návrhu')
p('Veřejná specifikace používá pouze obecné funkční potřeby: podporu porozumění řeči, krátké úkoly, přiměřenou zrakovou zátěž a možnost reagovat ukázáním. Neobsahuje fotografii zprávy, jméno, datum vyšetření, rok narození ani podrobná individuální měření. Není zdravotním záznamem konkrétní osoby.')
p('Uvažovaný profil zahrnuje potíže se soustředěním, potřebu jazykové podpory a kratší výdrž při čtení. Případná nepohoda může zůstat nevyjádřená. Krátká výdrž sama neurčuje příčinu; nehlášené obtíže ani odpověď „nevím“ neznamenají potvrzené pohodlí.')
table(['Funkční potřeba', 'Důsledek pro aplikaci'], [
('Mladší školní věk','Velké obrázky, jednoduché ovládání a přítomnost rodiče.'),
('Porozumění řeči','Krátké známé pokyny, poslech a možnost opakování.'),
('Kolísající výdrž','Dva krátké bloky, přestávka a snadné ukončení.'),
('Možné potíže s ostřením','Přiměřená práce na blízko; žádné vlastní předepisování oční léčby.'),
('Obtížné vyjádření nepohody','Volby dobře, pauza, konec a nevím bez nutnosti vysvětlení.')
],[6,11])
h('Zrakové souvislosti')
p('Při akomodačních potížích může být zraková ostrost dobrá, a přesto mohou být náročné úkoly na blízko. Akomodační exces je pojem pro nadměrné zapojení ostření nebo obtíž s jeho uvolněním. [S2] Aplikace tento stav nevyšetřuje ani neléčí. Z očního vyšetření nelze samostatně odvozovat deficit pozornosti, paměti či řeči.')

page('1 Kroky pro rodiče a pravidelný režim')
h('Co vyřešit v nejbližší době')
bullets([
'S pracovištěm upřesnit, co znamená RE, jaké konkrétní cviky předepsalo, jejich pomůcky, délku, četnost, důvody k přerušení a termín kontroly. Vyžádat také vlastní RightEye výstup, pokud existuje. Tato stránka jej neobsahuje.',
'S dětským očním lékařem ověřit stav refrakce a zda již bylo podle potřeby provedeno vyšetření v cykloplegii. Z dobrého visu nelze odvozovat, že další vyšetření ostření není potřebné. Léčbu ani brýle z fotografie nevolíme.',
'U klinického logopeda upřesnit porozumění jednoduchým větám, slovní zásobu a jazykové cíle. Při přetrvávajících obtížích pozornosti doma i ve škole konzultovat pediatra a dětského psychologa; podle dosavadních vyšetření také sluch. Oční nález jazykové potíže sám nevysvětluje. [S3]'
])
h('Praktický začátek doma')
p('Každý den zvolit přibližně 10 minut společné, klidné činnosti v době, kdy není unavený. Délky jsou návrhem snesitelného pilotu, ne ověřenou léčebnou dávkou. Není potřeba denně plnit všechny položky.')
table(['Kdy', 'Konkrétní činnost'],[
('Denně 5 až 7 minut','Společný příběh ze tří obrázků nebo knížky s velkými ilustracemi. Ptát se nejprve „Kdo?“ a „Co dělá?“. Ukázání je platná odpověď. Dospělý nabídne krátký správný jazykový vzor.'),
('Denně v běžné situaci','Jeden srozumitelný pokyn, potom klidné čekání přibližně 5 až 10 sekund. Např. „Dej lžíci na stůl.“ Dva kroky přidat až po zvládnutí jednoho. Nepřidávat najednou nová slova i více kroků.'),
('Třikrát týdně místo dalšího drilu','2 až 3 minuty opakování pořadí s figurkami či velkými kartami. Nebo pohybové „udělej jako já“ s jedním až dvěma kroky.'),
('Pravidelně během dne','Pohyb a hra venku, dostatek spánku a přestávky od čtení i displeje. Tablet má nahradit část jiné obrazovkové zábavy, nikoli navyšovat celkovou zátěž.')
],[4.4,12.6])
h('Týden s tabletem')
p('Pondělí: G1 paměť + G2 pokyny. Středa: G3 příběh + G1. Pátek: G2 + G4 hledání. Úterý, čtvrtek a víkend mohou zůstat bez aplikace. Jedna návštěva aplikace obsahuje dva bloky po nejvýše 2 minutách a mezi nimi alespoň minutu mimo obrazovku; pak konec. Příběh lze doříct ústně po odložení tabletu.')
p('Při pohodlném průběhu po dobu 1 až 2 týdnů lze s rodičem zvážit dva tříminutové bloky. Nezvyšovat zároveň délku a obtížnost. Pokud už má předepsaný zrakový program, uvedený rozvrh mu přizpůsobit a nesčítat oba tréninky mechanicky.')

page('2 Zraková pohoda a změny původního návrhu')
p('Tablet postavit stabilně, použít velké ovládání, příjemné osvětlení bez odlesků a takovou vzdálenost, aby se dítě nemuselo naklánět k drobným detailům. Přesné uspořádání přizpůsobit zařízení a doporučení vyšetřujícího. AAPOS doporučuje při delší práci přestávky s pohledem do dálky; náš kratší režim přerušuje zátěž dříve. Přestávka není cvičení rychlého přeostřování. [S4]')
p('Při rozmazání, dvojitém vidění, bolesti očí či hlavy hru ukončit. Při opakování obtíží konzultovat pracoviště před dalším tréninkem. Novou náhlou poruchu vidění nebo silnou bolest řešit neodkladně zdravotnicky. Běžný neúspěch ve hře není důvodem nutit dítě pokračovat.')
p('První týden krátce zaznamenat, kdy chce přestat a co rodič skutečně vidí: například mnutí očí, přivírání oka, přibližování k textu nebo ztrácení místa. Žádný jednotlivý projev neurčuje příčinu. Nabídnout pauzu před vyčerpáním; neprodlužovat čtení do potíží. Všímat si i pohodlí při poslechu příběhu a práci s reálnými předměty, bez diagnostických závěrů z tohoto porovnání.')
table(['Původní hra', 'Úprava pro nynější profil'],[
('Zvířátka na návštěvě','G1: zachovat. Velké statické domečky, pomalá sekvence, oddělené zkoušení ovládání a paměti.'),
('Nakrm správné zvíře','G2: vysoká priorita. Přirozené české pokyny, zkouška znalosti slov, samostatná obtížnost řeči a paměti.'),
('Lesní detektiv','G4: klidné hledání mezi 2 až 6 velkými obrázky. Vynechat husté scény, rychlé pohyby a časové závody.'),
('Semafor pro vláček','Odložit do druhé etapy. Statický signál, předvídatelné tempo, nejdřív ověření porozumění pravidlu.'),
('Třídírna pokladů','Druhá etapa. Pravidlo stále viditelné; změna výslovně oznámená. Nevynucovat rychlé přepínání.'),
('Co se stalo potom','G3: vysoká priorita, společně s rodičem. Tři velké scény a možnost ukázat místo mluvení.'),
('Cesta za pokladem','Druhá etapa. Malá přehledná plocha a nejvýše několik kroků, bez plynulého sledování pohybu.'),
('Rytmický parťák','Hlavně mimo displej s rodičem. Koordinace a pořadí, bez tvrzení o léčbě „spojováním hemisfér“.')
],[5.4,11.6])
p('Do této verze nezařazovat pencil push-ups, Brockův provázek, zakrývání oka, trénink s čočkami ani vynucené rychlé změny dálka–blízko. Zpráva neuvádí vergenční poruchu a neobsahuje předpis těchto postupů. Úprava velikosti nebo rozmazání obrázku na jednom displeji sama nevytváří skutečnou změnu optické vzdálenosti cíle; hry proto nejsou náhradou akomodačního tréninku.')
p('Kognitivní hry mohou zlepšovat procvičované schopnosti, ale širší přenos zůstává proměnlivý. [S5, S6] U jazykových obtíží je vhodné spojit hru s cíli logopeda. [S3, S7, S8] Z herního výkonu se nebude odhadovat ADHD, dysfázie ani stav očí.')

page('3 Společné požadavky na první verzi')
p('První verze je česká webová aplikace pro tablet s rodičem. Obsahuje G1 až G4, rodičovské nastavení, krátký průvodce mimo displej a lokální historii. Každá hra má jediný srozumitelný cíl. Ilustrace a slovník zůstávají konzistentní napříč hrami.')
h('Rozhraní a obsah')
bullets([
'Primární ovládání klepnutím. Přetahování není povinné. Aktivní plochy nejméně 64 × 64 CSS px, mezery nejméně 16 px, obvykle 2 až 6 objektů; skutečnou čitelnost ověřit na cílovém tabletu, CSS pixely nejsou fyzické milimetry.',
'Velké jednoduché ilustrace bez pozadí, stabilní rozmístění v průběhu úlohy. Barvu vždy doplnit tvarem či symbolem. Žádné blikání, paralaxa, automatický zoom nebo zbytečná kamera v pohybu.',
'V první verzi české namluvené věty přibalené k aplikaci, nikoli povinná online syntéza. Rodič může instrukci přečíst. Bez hudby na pozadí. Tlačítko „Ještě jednou“ je dostupné; opakování se zapisuje jako podpora.',
'Dítě nemusí číst, správně vyslovovat ani vysvětlovat pravidlo slovy. Porozumění pravidlu ukáže na názorných pokusech. Výslovnost ani řeč nehodnotí automatický rozpoznávač.',
'„Pauza“ a „Končíme“ jsou dostupné stále. Pochvala za pokus a dokončení; bez životů, skóre inteligence, denních sérií, reklam, nákupů a porovnávání s jinými dětmi.'
])
h('Průběh návštěvy')
p('Rodič zvolí dvě hry. Krátká kontrola pohodlí → názorná ukázka → dva pokusy s podporou → herní blok → přestávka mimo displej → druhý blok → krátké shrnutí pro rodiče. Když pravidlo není jasné ani po opětovné ukázce, hra nabídne společné zkoušení nebo konec; výkon se neoznačuje za nepozornost.')
p('Výchozí blok trvá nejvýše 120 sekund viditelného času v aplikaci včetně instrukce a zpětné vazby. Na hranici limitu zastavit scénu, rozehraný pokus označit jako nedokončený a nabídnout přestávku, nikoli chybu. Během přestávky zhasne herní plocha, dítě se dívá mimo displej; pokračování nejdřív po 60 sekundách a po potvrzení rodičem. Za jednu návštěvu nejvýše 240 sekund a dva bloky.')
p('Skrytí záložky zastaví zvuk a vstupy. Po návratu se přerušený pokus nezapočítá jako chyba a spustí se nový po potvrzení. Pauza ani obnovení stránky neresetují už vyčerpaný čas návštěvy. Rodičovská změna na 180 sekund na blok nastaví limit návštěvy na 360 sekund; dítě délku samo nemění.')
p('Do limitu patří viditelný čas ukázky, nácviku, pokusu a zpětné vazby. Rodičovské nastavení a prázdná plocha pauzy se nepočítají. Novou návštěvu spouští rodič; výchozí plán dovoluje jednu návštěvu v daný den. Opuštění a nové otevření aplikace nesmí samo vytvořit další návštěvu.')
h('Zpětná vazba')
p('Po správném pokusu krátké „Povedlo se“. Při chybě neutrální „Zkusíme to spolu“ a ukázka. Žádný nepříjemný zvuk. Dlouhé čekání neukončuje pokus jako chybný; po přibližně 15 sekundách nabídnout rodiči pomoc nebo přeskočení. Celkový časový limit bloku však platí dál.')
p('Dítě může ukázat na obrázek „Je mi dobře“, „Chci pauzu“, „Chci skončit“ nebo „Nevím“. Nemusí důvod popsat ani potvrdit bolest. „Nevím“ zůstává nejasným údajem; žádost o pauzu sama neznamená bolest. Volbu dítěte ukládat odděleně od pozorování rodiče. Přestávky platí i při nehlášených obtížích; jejich využití nemá vliv na odměnu.')

page('4 Hra G1 Zvířátka na návštěvě')
p('Cíl: procvičit zapamatování prostorové sekvence s minimální jazykovou zátěží. Hra netestuje kvalitu očních pohybů. Použít stále stejné velké domečky a jedno snadno rozeznatelné zvířátko.')
h('Přesný průběh pokusu')
bullets([
'Na ploše jsou čtyři domečky v mřížce 2 × 2. V ukázce se dítě nejprve učí pouze klepnout na zvýrazněný domeček; tyto pokusy se nehodnotí jako paměť.',
'Instrukce „Podívej, kam šlo zvířátko.“ Zvířátko se bez přesouvací animace postupně objeví v domečcích. Každý cíl je vidět 1 500 ms, mezi cíli je 500 ms klidu. Domečky zůstávají na místě.',
'Po sekvenci zvířátko zmizí, ozve se „Teď ty“. Vstup začne až po skončení instrukce. Dítě klepne na domečky ve správném pořadí. Během ukázky se dotyky nehodnotí.',
'Správnost se vyhodnotí po celé odpovědi, ne okamžitě po první chybě. Při chybě ukázat správné pořadí a nabídnout jeden společný pokus se stejnou sekvencí; původní výsledek se nepřepisuje.',
'Tlačítko zopakování předvede celou sekvenci znovu a zapíše replay. Stejná sekvence se neopakuje v bezprostředně dalším samostatném pokusu.'
])
table(['Úroveň', 'Parametry'],[
('Nácvik','1 cíl ze 4; společně, bez postupového skóre'),
('L1 výchozí','2 cíle ze 4; bez opakování domečku v sekvenci'),
('L2','3 cíle ze stejných 4; stejná rychlost a rozvržení'),
('L3','4 cíle ze stejných 4; stejná rychlost a rozvržení')
],[4,13])
p('Délku sekvence nezvyšovat současně se zkrácením zobrazení. V první verzi je tempo pevné. Při obtížích lze sestoupit do nácviku; není stanoven cílový výkon podle věku.')
h('Měření a přenos')
p('Ukládat délku sekvence, identifikátory cílů, odpověď, správnost prvního pokusu, podporu, přerušení a aktivní dobu. Přehled ukazuje například „2 kroky, samostatně 4 z 5 pokusů“. Mimo tablet rodič postaví čtyři misky nebo velké karty a předvede obdobné pořadí rukou. Nové sekvence a reálné předměty brání pouhému zapamatování konkrétní herní sady.')
h('Přejímka G1')
p('Generátor dodržuje délku i zákaz opakování v jedné sekvenci. Dítě nemůže reagovat během prezentace. Opakování vždy označí pokus jako podpořený. Změna velikosti obrazovky nebo přerušení sekvence pokus zneplatní bez penalizace. Vyhodnocuje se celé pořadí, nikoli jen počet zasažených domečků.')

page('5 Hra G2 Pomocník se zvířátky')
p('Cíl: procvičovat porozumění známému krátkému pokynu a později pořadí dvou pokynů. Tento modul má největší vazbu na podezření na jazykové obtíže; přesný obsah má před použitím zkontrolovat rodič a podle možností logoped.')
h('Ovládání a ověření slov')
p('Nejprve společně pojmenovat předměty a postavy a nechat dítě ukázat, co už zná. Neznámé slovo se učí mimo hodnocené pokusy. Obsluha je „klepni na předmět, klepni na příjemce“. Vybraný předmět má zřetelný rámeček a lze výběr zrušit; tažení není nutné. Dítě má k dispozici celý čas bloku.')
table(['Režim', 'Příklad a pravidlo'],[
('A Jeden krok','„Dej jablko medvědovi.“ Dva známé předměty a dvě známé postavy. V každém pokusu změnit jen kombinaci příjemce a předmětu.'),
('B Dva kroky','„Dej jablko medvědovi. Potom dej mrkev králíkovi.“ Stejná slovní zásoba a stejná scéna jako A. Pokyny zazní v jednom celku, dítě pak provede oba.'),
('C Jedna jazyková změna','Volitelně až po logopedickém výběru cíle. Např. „Dej míč do krabice.“ Začít znovu jedním krokem. Novou předložku netrénovat současně se dvěma kroky.')
],[4.2,12.8])
h('Podpora bez zkreslení výsledku')
p('Nácvik může obsahovat obrázkový proužek ukazující správný předmět a příjemce. V samostatném režimu tato nápověda zpočátku vidět není, jinak by šlo splnit úlohu bez porozumění řeči. Po stisknutí nápovědy se proužek objeví a výsledek je podpořený. Rodič může zadání zopakovat či zjednodušit; obě možnosti se zaznamenají samostatně.')
p('Příklad: dítě po první větě vybere jiný předmět. Aplikace nabídne zopakování, případně obrázkový vzor. Netvrdí, zda příčinou byla paměť, slovník, porozumění, zrak nebo pozornost. Po dvou neúspěšných nácvicích rodič zvolí jednodušší úkol či společnou hru.')
h('Postup a obsah')
p('Režim B se nabídne až po zvládnutí A podle pravidla postupu a potvrzení rodičem; význam slova „potom“ předem ověřit v nácviku. Režim C má vlastní úroveň i historii. Obsah pro první vydání: alespoň 12 jednoznačných kombinací A, 12 dvoukrokových kombinací B a dvě ukázky každého režimu. Generátor nepoužije neznámé slovo ani neproveditelný přesun.')
h('Měření a přejímka')
p('Zapisovat správný předmět, příjemce a pořadí jednotlivě; úspěch celého pokusu pouze při správném dokončení všech kroků. Odlišit vlastní volbu, audio replay, obrázkovou nápovědu a pomoc rodiče. Mimo tablet použít jiný známý předmět v domácím pokynu. Přejímka musí ověřit, že A a B mění počet kroků bez změny jazykové obtížnosti a že aktivní nápověda nikdy nevypadá jako samostatný úspěch.')

page('6 Hra G3 Příběh ve třech obrázcích')
p('Cíl: společně rozvíjet posloupnost, porozumění ději a vyjadřování. Příběhy jsou krátké a každodenní. Dospělý poskytuje jazykový vzor, nikoli zkoušení pod tlakem. Výzkum strukturované jazykové intervence podporuje tento směr, ale účinnost konkrétní aplikace tím není ověřena. [S8]')
h('Pokus a konkrétní obsah')
p('Příklad „Rozlitá voda“: 1. dítě převrhne sklenici; 2. voda je na stole a dítě bere utěrku; 3. dítě stůl utře. Obrázky se liší velkou jasnou akcí, ne drobnými detaily. Alternativně „Mokré boty“: začne pršet → obujeme holínky → jdeme ven do louže.')
bullets([
'Rodič s dítětem nejprve pojmenuje postavy a předměty. Následuje zcela nový příběh ze stejné známé slovní zásoby.',
'Tři velké karty jsou v náhodném pořadí. Dítě klepne na kartu a potom na očíslované místo; číslo doprovází jasná poziční značka. Přetahování je pouze alternativní.',
'Po seřazení rodič položí jednu otázku, např. „Co se stalo první?“ Dítě může odpovědět ukázáním. Až potom může následovat „Co udělá dál?“ nebo jednoduché „Proč?“.',
'Rodič naváže krátkou větou: „Voda se rozlila. Vezmeme utěrku.“ Dítě může zopakovat nebo doplnit, ale mluvení se nevynucuje.',
'Při konci časového bloku se obrazovka odloží a povídání může pokračovat ústně. Rozpracovaný digitální pokus se neoznačí za chybu.'
])
h('Obtížnost')
p('L1: dvě zřetelně navazující scény a otázka kdo/co. L2: tři scény a otázka na pořadí. L3: tři scény a jedna příčina nebo jednoduchá předpověď, zvolená logopedem či rodičem. V první verzi nepřidávat čtvrtou scénu ani nové složité souvětí. Jazykovou obtížnost nastavuje rodič; aplikace ji automaticky nezvyšuje podle pořadí karet.')
h('Měření')
p('Automaticky se hodnotí pouze pořadí karet. Jazykový projev rodič označí jako „ukázal“, „odpověděl slovem“, „odpověděl větou“, „s pomocí“ nebo „nehodnoceno“. Tyto volby nejsou stupnicí inteligence ani automatickým měřítkem zlepšení. Nenahrávat hlas nebo video.')
h('Přejímka a přenos')
p('Pro vydání připravit nejméně 12 originálních příběhů, každý se třemi scénami, správným pořadím, jednoduchou otázkou, přijatelnými odpověďmi a krátkým jazykovým vzorem. Kde jsou možné dvě logické posloupnosti, uznat obě nebo obsah upravit před vydáním. Rodičův záznam se nesmí doplnit automaticky. Přenos sledovat na jiné knížce nebo na vyprávění o skutečné události, kterou aplikace nepoužila.')

page('7 Hra G4 Klidný detektiv')
p('Cíl: procvičit výběr cíle mezi několika jasně oddělenými obrázky. Hra slouží jako klidné hledání; nebude se vydávat za léčbu sakád nebo vyšetření trvalé pozornosti.')
h('Přesný průběh')
bullets([
'V horní části je stále vidět velký vzor cíle. Instrukce „Najdi stejný obrázek.“ Ve spodní části se zobrazí dva až šest objektů. Vzor odstraňuje požadavek pamatovat si zadání.',
'Na každé ploše existuje právě jeden správný cíl. Obrázky stojí na místě až do odpovědi. Nesbírají se pohyblivé objekty a obrázky neblikají.',
'Dítě klepne. Po chybě zůstane scéna stejná a rodič může pomoci systematickým prohlédnutím. První volba se zachová v záznamu, oprava ji nepřepíše.',
'Není zobrazen odpočet ani časový žebříček. Pomalá správná volba má stejnou odměnu jako rychlá. Nepoužívat velmi podobné detaily typu odlišný počet drobných teček.'
])
table(['Úroveň', 'Parametry'],[
('L1','2 velké objekty; cílové zvíře a zřetelně jiný předmět'),
('L2','4 objekty na pevné mřížce 2 × 2; stejná velikost a styl'),
('L3','6 objektů na mřížce 3 × 2; stejné snadno odlišitelné tvary')
],[4,13])
p('Nezmenšovat cíle kvůli počtu objektů. Pokud cílové zařízení neposkytne potřebné místo, zůstat na nižší úrovni. Pozice cíle se mezi pokusy vyvažují; žádná pozice není trvale správná. V jedné scéně se pozice nemění.')
h('Obsah a měření')
p('Připravit alespoň 12 jednoznačných ikon a pro každou úroveň generátor scén s právě jedním cílem. Ukládat první volbu, opravu, pomoc, počet objektů a čas odpovědi po skončení instrukce. Čas je popisný údaj, nikdy podklad pro samostatné tvrzení „lepší soustředění“. Přerušení, změna orientace a pauza zneplatní časové srovnání daného pokusu.')
h('Přejímka a přenos')
p('Automatická kontrola musí zachytit chybějící i duplicitní cíl. Vzor je stále viditelný. Všechny ikony splňují minimální velikost a žádná není oříznutá na podporovaných rozměrech. Každou polohu lze ovládat dotykem i klávesnicí. Mimo tablet dítě hledá jeden předmět mezi třemi až pěti skutečnými předměty na stole; změna se hodnotí odděleně od hry.')
h('Zařazení do pilotu')
p('Začít pouze jedním krátkým blokem v pátek. Pokud je hledání únavné, vynechat ho a použít příběh s rodičem nebo aktivitu mimo displej. Nízký výkon nepovažovat za důkaz poruchy očních pohybů; pro individuální zrakový program je potřeba skutečný výstup pracoviště.')

page('8 Další hry po ověření první verze')
p('Tyto moduly patří do druhé etapy. Přidávat je podle zájmu dítěte a konkrétního cíle, nikoli proto, aby každý den trénovalo všechny kognitivní oblasti. Přesná časování níže jsou návrhy implementace a musejí projít zkouškou s dospělým a následně krátkým pilotem.')
h('G5 Semafor')
p('Statický kruh s obrázkem ruky znamená klepnout na velké tlačítko, osmiúhelník s dlaní znamená počkat. První blok šest ukázkových pokusů po 3 sekundách s rovnoměrným zastoupením. Po zvládnutí pravidla deset pokusů se sedmi „klepni“ a třemi „počkej“, mezi nimi 1 sekunda prázdné plochy. Každý signál trvá celé 3 sekundy; nepřerušovat ho rychlou správnou reakcí. Vyhodnocovat chybnou reakci na stop a chybějící reakci zvlášť. Doba reakce není cílem. Nejasné pravidlo vede zpět do nácviku, ne k diagnóze impulzivity.')
h('G6 Třídírna')
p('Jeden předmět a dvě nádoby. Pět pokusů třídění podle barvy, potom názorné oznámení a ukázka nového pravidla podle tvaru. Aktuální pravidlo je stále vidět. Začít bloky jednoho pravidla; teprve později střídat po pěti úlohách. Změnu nikdy neprovést nečekaně. Hodnotit správnost a potřebnou nápovědu. Kombinace vlastností musí zabránit tomu, aby staré i nové pravidlo vždy dávaly stejnou odpověď.')
h('G7 Cesta za pokladem')
p('Statická mřížka 3 × 3, start a cíl. Dítě zvolí dva směrové kroky velkými tlačítky, potom stiskne „Zkusit“. Postava se po jednotlivých krocích objeví na dalším poli bez plynulého posunu kamery. L1 bez překážky, L2 jedna překážka při stejné délce, L3 tři kroky. Přijmout každou platnou cestu; mimo mapu nepokračovat. Hodnotit splnění cíle a samostatné opravy, nikoli rychlost.')
h('G8 Rytmický parťák')
p('Přednostně s rodičem mimo displej: dospělý předvede dva údery do stehen nebo dvě tlesknutí, dítě zopakuje. Postup na tři prvky až při pohodlném zvládání. Tablet může jednou přehrát ukázku, pak se odloží. Nepředepsat dominanci ruky ani nucené přeučování laterality. Digitální modul neměří synchronizaci mozkových hemisfér a nepřisuzuje jí léčebný účinek. [S10]')
h('Co do druhé etapy nepřechází automaticky')
p('Vlastní vergenční, akomodační nebo RightEye terapie není součástí těchto her. Případné přidání konkrétního očního cvičení je samostatné zadání založené na předpisu odborníka, potřebných pomůckách a ověření proveditelnosti. Důkazy pro léčbu jiné zrakové poruchy se nepřenášejí na tento profil. [S9]')

page('9 Adaptace a měření přínosu')
h('Jednoznačné pravidlo obtížnosti')
p('G1, G2 A/B a G4 mají oddělené úrovně. Nácvik, přeskočení, přerušené a technicky neplatné pokusy se do postupového okna nepočítají. Dokončené pokusy s nápovědou se počítají, ale nejsou samostatným úspěchem. Okno tvoří posledních 10 dokončených prvních pokusů na stejné úrovni a ve stejné verzi obsahu.')
p('Nabídnout rodiči zvýšení o jeden stupeň pouze tehdy, pokud okno zahrnuje nejméně dvě různá data, alespoň 8 z 10 pokusů je správně bez pomoci a nejvýše jeden pokus použil podporu. Rodič musí zároveň potvrdit pohodlí dítěte. Ke zvýšení dojde až v příští návštěvě; nezmění se současně tempo, délka bloku ani jiný parametr. Po změně začíná nové postupové okno. Toto je produktové pravidlo, nikoli klinický práh.')
p('Po třech chybných prvních pokusech za sebou se zobrazí pauza a nabídka ověřit pravidlo, pomoci nebo skončit. Úroveň se automaticky nesnižuje na základě únavy nebo pomalosti. Rodič může kdykoli zvolit snazší variantu; změna a důvod se zapíší. U G3 rodič volí jazykovou obtížnost ručně.')
h('Co sledovat mimo aplikaci')
p('Před prvním týdnem během tří běžných dnů zaznamenat výchozí stav. Potom jednou týdně ve srovnatelné situaci udělat krátké pozorování bez herních odměn. Neprodlužovat kvůli měření domácí úkoly. Výsledky jsou individuální popis, bez věkových percentilů.')
table(['Ukazatel', 'Způsob záznamu'],[
('Běžný pokyn','Ze tří přirozených příležitostí s dobře známými slovy kolik zvládl samostatně. Jeden a dva kroky vést odděleně.'),
('Samostatná práce','U podobně náročné zhruba pětiminutové činnosti počet připomenutí a zda ji dokončil. Poznamenat únavu či změnu zadání.'),
('Příběh','Nové tři obrázky: rozpoznal děj, určil pořadí, potřeboval podporu. Bez známkování výslovnosti.'),
('Pohoda','Zvlášť volba dítěte (dobře / pauza / konec / nevím) a pozorování rodiče. Nehlášená bolest se nepřevádí na potvrzené pohodlí. Zapsat i důvod ukončení, je-li znám.')
],[4.4,12.6])
p('Po dvou týdnech vyhodnotit hlavně snášenlivost a zájem. Po čtyřech týdnech porovnat běžné situace, případně požádat učitele o stručné stejné pozorování. Když roste pouze skóre hry, evidujeme naučení hry, nikoli prokázané zlepšení koncentrace. Pokud není praktický přínos nebo vzniká odpor, zkrátit či změnit činnost a probrat cíle s odborníkem.')

page('10 Technický návrh a ochrana dat')
p('Doporučené provedení: React + TypeScript + Vite, responzivní web s možností PWA a offline provozu po prvním načtení. Stabilní verze knihoven zvolit a ověřit při implementaci. Pro první verzi stačí lokální úložiště IndexedDB; serverový účet není potřeba. Instalace PWA zůstává volitelná a web musí fungovat i v běžném prohlížeči.')
h('Oddělené části aplikace')
bullets([
'SessionController: řídí návštěvu, viditelný čas, dva bloky, přestávku, pauzu a ukončení. Je společný všem hrám; přepnutí hry neobchází limit.',
'GameEngine: čisté vyhodnocovací funkce, generování obsahu s uloženým seedem, oddělené úrovně a pravidla podpory. Žádná diagnostická inference.',
'ContentPack: verzované identifikátory obrázků a nahrávek, slovník, obtížnost, správné odpovědi a alternativa mimo displej. V MVP ručně kontrolovaný obsah, bez dynamicky generovaných pokynů.',
'ParentPanel: nastavení, přehled, příznaky zadané rodičem, export, import JSON a odstranění dat. Vstup přes podržení a dospělý text; nejde o zabezpečení před jiným uživatelem zařízení.'
])
h('Datové typy')
p('Session: id, schemaVersion, contentVersion, localDate, startedAt, endedAt, status, gameIds, visibleMs, blockLimitMs, sessionLimitMs, breakCompleted, comfortBefore, comfortAfter, stopReason. Comfort obsahuje odděleně childChoice (okay, break, end, unsure, notAsked) a parentObservedSigns[]. Povolené stavy návštěvy: inProgress, completed, stopped, interrupted.')
p('Trial: id, sessionId, gameId, mode, level, contentId, seed, target, firstResponse, completed, correctFirst, supportTypes[], replayCount, responseMs nebo null, invalidReason nebo null. supportTypes = audioReplay, visualCue, parentHelp, simplifiedInstruction. Nácvik nese practice=true. G3 navíc samostatný volitelný parentLanguageObservation.')
p('Observation: id, localDate, context, measure, opportunities, independentSuccesses, prompts, comfort a krátká volitelná poznámka. Žádný údaj nepřevádět na „procento soustředění“. Chybějící měření zobrazit „nezaznamenáno“, nikdy nulu. Export JSON uchovává úplné struktury; CSV je zploštělý přehled návštěv a pozorování.')
h('Soukromí a technická omezení')
p('Kód ani veřejné soubory neobsahují jméno, rok narození, fotografii zprávy ani diagnózu dítěte. Výchozí profil je „Hráč“. Bez analytických služeb, reklam, externích fontů, mikrofonu, kamery a cloudového přenosu výsledků. Obrázky a zvuk jsou součástí aplikace. Export provádí rodič vědomým stažením.')
p('Lokální data nejsou automaticky šifrovaná ani zálohovaná. Smazání úložiště prohlížeče je může odstranit; upozornit v rodičovském panelu a nabídnout export. Zabezpečení zařízení řeší jeho uživatelský účet. Pro veřejné nasazení znovu prověřit hosting, skutečné síťové požadavky a práci s dětskými údaji; tento dokument a zdravotní podklady se nepublikují.')

page('11 Přejímací podmínky a pořadí vývoje')
table(['ID', 'Ověřitelná podmínka'],[
('AC01','Dvě ukázky a ovládání fungují bez čtení. Povinné jazykové odpovědi nejsou v žádné hře.'),
('AC02','Pauza, skrytí stránky a návrat zastaví zvuk i odpovědní čas; rozpracovaný pokus se nestane chybou.'),
('AC03','Po 120 sekundách viditelné aktivity nastane přestávka. Před 60 sekundami nelze pokračovat. Po dvou blocích se návštěva ukončí.'),
('AC04','Obnovení stránky obnoví stav návštěvy a vyčerpaný limit. Volba jiné hry limit nevynuluje. Režim 180 sekund má limit 360 sekund.'),
('AC05','Nápověda a pomoc rodiče nikdy nejsou vykázány jako správný samostatný první pokus.'),
('AC06','Postupový algoritmus odmítne 8 úspěchů z jediného dne, neúplné okno i smíšené úrovně; správné podmínky pouze nabídnou postup rodiči.'),
('AC07','G1 ověří pořadí, G2 kroky a příjemce, G3 připouští platné alternativy, G4 má právě jeden cíl.'),
('AC08','Tablet v orientaci na šířku i výšku zůstává čitelný; nic není oříznuté. Dotykové plochy mají alespoň 64 px. Změna rozvržení přeruší aktivní pokus.'),
('AC09','Po prvním načtení fungují hry i přibalené audio offline. Chybějící audio má text pro rodiče a ukázku, ne neovladatelnou obrazovku.'),
('AC10','Export a opětovný import zachovají počty, podporu a nevyplněné údaje; po potvrzeném smazání nejsou záznamy v úložišti.'),
('AC11','Síťová kontrola neukáže odchozí zdravotní či herní výsledky. Ve veřejném balíčku nejsou osobní údaje z tohoto zadání.'),
('AC12','Příznak bolest/rozmazání/dvojité vidění ukončí návštěvu. Nenabídne další pokus ani odměnu za překonání obtíží.'),
('AC13','„Nevím“ není „v pohodě“; rodičovské pozorování nepřepisuje odpověď dítěte. Pauza i konec fungují bez uvedení příčiny a bez penalizace.')
],[1.8,15.2])
h('Implementační postup')
p('1. Společné rozhraní a řízení času, pauzy, instrukcí a rodičovského nastavení. 2. G1 a G2 s obsahovými příklady. 3. G3 a G4, přehled a export. 4. Kontrola všech her dospělým, jazyková kontrola, test na skutečném tabletu. 5. Krátký rodinný pilot a úpravy podle pohodlí. Teprve potom rozšiřovat o G5 až G8.')
p('Testy mají ověřovat rizikové chování: časové limity, přerušení, podporu, generování jednoznačných úloh, import a úplné smazání. Dále ručně ověřit Safari na iPadu a Chrome na Android tabletu, zvuk po uživatelském dotyku a offline režim. Dokud neproběhne fyzický test cílového zařízení, uvádět jeho kompatibilitu jako neověřenou.')

page('12 Otevřené údaje a podklady')
h('Co doplnit při dalším kroku')
bullets([
'Při navrhování konkrétního domácího režimu ověřit výdrž u čtení a případné obtíže. Doplnit pozorované projevy a okolnosti ukončení čtení; případně aktuální doporučení vyšetřujícího pracoviště.',
'Význam RE a konkrétní předepsaný postup; samostatný RightEye protokol včetně věkových norem a interpretace pracoviště.',
'Aktuální logopedické cíle, slova kterým bezpečně rozumí a počet kroků, který zvládá ve známé situaci.',
'Cílový tablet, prohlížeč, velikost displeje, oblíbené téma a přibližná dosavadní obrazovková zátěž.'
])
p('Tyto údaje zpřesní obsah a režim. Vývoj obecné klidné aplikace mohou doprovázet průběžně; cílený oční trénink bez svého předpisu do aplikace nevkládáme. Volba délky pilotu, herních úrovní a postupových prahů je součástí návrhu produktu, nikoli závěrem očního testu.')
h('Zdroje a rozsah jejich použití')
p('S1 Neveřejné vstupní podklady byly při přípravě veřejné specifikace zobecněny. Individuální zdravotní zpráva ani její identifikátory nejsou součástí veřejné dokumentace.')
sources=[
('S2 SUNY University Eye Center — Accommodative Problems','https://www.universityeyecenter.org/focusing-problemsaccommodative-problems/','Vysvětlení poruch ostření; ne individuální léčebný postup.'),
('S3 NIDCD — Developmental Language Disorder','https://www.nidcd.nih.gov/health/developmental-language-disorder','Porozumění pokynům, vyšetření řeči a cílená jazyková podpora.'),
('S4 AAPOS — Screen Time and Online Learning','https://aapos.org/glossary/screen-time-and-online-learning','Zraková hygiena a přestávky. Kratší pilotní bloky jsou vlastní návrh.'),
('S5 Westwood a kol 2023 — Metaanalýza kognitivního tréninku','https://www.nature.com/articles/s41380-023-02000-7','36 RCT u ADHD: zlepšení pracovní paměti, omezenější klinický přenos.'),
('S6 Yan a kol 2026 — Computerized executive function training','https://doi.org/10.1016/j.jad.2025.120730','Příznivější závěry širší metaanalýzy, většinou střední či vysoké riziko zkreslení. Výsledky výzkumu nejsou jednotné.'),
('S7 Henry a kol 2022 — Working memory intervention in DLD','https://www.mdpi.com/2076-3425/12/5/642','47 dětí 6–10 let; zlepšení pracovní paměti a porozumění větám. Výsledek jiné intervence není ověřením našich her.'),
('S8 Petersen a kol 2026 — Narrative and expository language intervention','https://pmc.ncbi.nlm.nih.gov/articles/PMC13114585/','Podpora strukturovaného jazykového nácviku; nepřenášet velikost účinku přímo na domácí aplikaci.'),
('S9 CITT ART 2021 — Therapy and attention','https://pmc.ncbi.nlm.nih.gov/articles/PMC8639028/','Léčba jiné diagnózy, konvergenční insuficience, neprokázala oproti placebu zlepšení pozornosti; není studií tohoto akomodačního nálezu.'),
('S10 Ferrero a kol 2017 — Crossed laterality meta analysis','https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0183618','Nepodporuje spolehlivou souvislost zkřížené laterality se školním výkonem či inteligencí. Hry nemají slibovat nápravu hemisfér.')
]
for label,url,explanation in sources:
    link(label,url)
    para=p(explanation); para.paragraph_format.space_after=Pt(3)
    for r in para.runs:r.font.size=Pt(9)

doc.save(OUT/'Domaci_plan_a_zadani_webovych_her.docx')

config={
 'schemaVersion':1,'profileName':'Hráč','locale':'cs-CZ',
 'purpose':'cognitive_and_language_practice','clinicalTreatment':False,
 'session':{'blockVisibleLimitMs':120000,'blocksPerSession':2,'sessionVisibleLimitMs':240000,'breakMinMs':60000,'parentConfirmResume':True,'suggestedSessionsPerWeek':3,'maxPlannedSessionsPerDay':1,'parentOptionalBlockMs':180000,'parentOptionalSessionMs':360000,'persistConsumedTime':True},
 'ui':{'minTouchTargetCssPx':64,'minGapCssPx':16,'backgroundMusic':False,'flashing':False,'forcedReading':False,'requiredSpeech':False,'dragRequired':False,'onboardingTrials':2},
 'comfort':{'childChoices':['okay','break','end','unsure','notAsked'],'labelsCs':['Je mi dobře','Chci pauzu','Chci skončit','Nevím','Neptali jsme se'],'pictureChoice':True,'reasonRequired':False,'parentObservationSeparate':True,'unknownIsComfortable':False,'breakRequestMeansPain':False,'scheduledBreaksEvenWithoutComplaints':True},
 'adaptation':{'windowCompletedTrials':10,'minDistinctLocalDates':2,'minCorrectUnassisted':8,'maxSupportedTrials':1,'parentConfirmation':True,'applyNextSession':True,'oneParameterAtATime':True,'consecutiveWrongOfferPause':3,'languageStoryLevelAutoAdvance':False},
 'games':{
  'G1':{'levels':[{'id':'L1','grid':[2,2],'sequenceLength':2},{'id':'L2','grid':[2,2],'sequenceLength':3},{'id':'L3','grid':[2,2],'sequenceLength':4}],'showTargetMs':1500,'gapMs':500,'repeatWithinSequence':False},
  'G2':{'initialMode':'A','modes':{'A':{'steps':1},'B':{'steps':2},'C':{'steps':1,'parentSelectedLanguageTarget':True}},'keepVocabularyAcrossAandB':True},
  'G3':{'levels':{'L1':{'scenes':2},'L2':{'scenes':3},'L3':{'scenes':3,'parentSelectedQuestion':True}},'speechRecording':False,'autoLanguageScoring':False},
  'G4':{'objectsByLevel':[2,4,6],'targetVisible':True,'responseDeadlineMs':None,'exactlyOneTarget':True}
 },
 'privacy':{'storage':'IndexedDB','accounts':False,'telemetry':False,'microphone':False,'camera':False,'export':['JSON','CSV'],'medicalReportBundled':False},
 'notEstablished':['clinicalEfficacy','deviceCompatibility','individualOptimalDose']
}
(OUT/'vychozi_nastaveni.json').write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf-8')
examples={
 'contentVersion':'0.1-examples','status':'examples_for_implementation_not_complete_content_pack',
 'G1':{'contentId':'memory-example-01','level':'L1','target':['house-1','house-3']},
 'G2':[
 {'contentId':'instruction-a-01','mode':'A','audioText':'Dej jablko medvědovi.','knownWordsRequired':['jablko','medvěd','dát'],'targets':[{'item':'apple','recipient':'bear'}]},
 {'contentId':'instruction-b-01','mode':'B','audioText':'Dej jablko medvědovi. Potom dej mrkev králíkovi.','knownWordsRequired':['jablko','medvěd','mrkev','králík','dát','potom'],'targets':[{'item':'apple','recipient':'bear'},{'item':'carrot','recipient':'rabbit'}]}
 ],
 'G3':{'contentId':'story-water-01','scenes':[{'id':'spill','description':'Dítě převrhne sklenici vody.'},{'id':'cloth','description':'Voda je na stole a dítě bere utěrku.'},{'id':'wipe','description':'Dítě utře stůl.'}],'acceptedOrders':[['spill','cloth','wipe']],'questions':[{'text':'Co se stalo první?','acceptableResponse':'Ukáže převrženou sklenici nebo popíše rozlití vody.'}],'parentModel':'Voda se rozlila. Vezmeme utěrku. Stůl utřeme.'},
 'G4':{'contentId':'search-example-01','level':'L1','targetIcon':'fox','options':['fox','cup'],'targetCount':1}
}
(OUT/'obsahove_priklady.json').write_text(json.dumps(examples,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created',OUT/'Domaci_plan_a_zadani_webovych_her.docx')
print('Created 2 implementation JSON files')
